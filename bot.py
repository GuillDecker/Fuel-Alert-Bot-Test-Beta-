#!/usr/bin/env python3
"""Fuel Alert Bot — prix maximum des carburants au Luxembourg et en Belgique.

Sources :
  - carbu.com/luxembourg : diesel, SP95, SP98 (Luxembourg)
  - petrol.lu            : mazout de chauffage (Luxembourg)
  - SPF Économie         : diesel, SP95, SP98, mazout (Belgique)
  - prix-carburants.gouv.fr : prix typique (médiane) hors autoroute (France)
  - Bulletin pétrolier UE   : mazout, moyenne nationale hebdomadaire (France)

La France est affichée à titre de comparaison : elle ne déclenche jamais de notification.
Chaque source est indépendante : si l'une tombe, les autres continuent.
"""

import csv
import io
import json
import os
import re
import smtplib
import ssl
import statistics
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta, timezone
from decimal import ROUND_DOWN, Decimal
from email.message import EmailMessage
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from bs4 import BeautifulSoup

# ─── Configuration ──────────────────────────────────────────────────
TZ = ZoneInfo("Europe/Luxembourg")
BASE_DIR = Path(__file__).resolve().parent
CACHE_FILE = BASE_DIR / "prices_cache.json"
HISTORY_FILE = BASE_DIR / "history.csv"

NTFY_URL = os.environ.get("NTFY_URL", "").strip()
NTFY_ALERT_URL = os.environ.get("NTFY_ALERT_URL", "").strip()
MAIL_ADDRESS = os.environ.get("MAIL_ADDRESS", "").strip()
MAIL_APP_PASSWORD = os.environ.get("MAIL_APP_PASSWORD", "").replace(" ", "").strip()

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                  "AppleWebKit/537.36 (KHTML, like Gecko) "
                  "Chrome/124.0.0.0 Safari/537.36"
}
REQUEST_TIMEOUT = (10, 20)   # (connexion, lecture) en secondes
REQUEST_ATTEMPTS = 2
RETRY_DELAY = 5

PRICE_MIN = Decimal("0.50")  # garde-fou anti-erreur de lecture
PRICE_MAX = Decimal("10.50")

FAIL_THRESHOLD = 12          # 12 runs consécutifs ≈ 1 h → alerte
NTFY_FAIL_THRESHOLD = 3      # 3 notifs refusées d'affilée ≈ 15 min → alerte
REMINDER_HOURS = 24          # rappel si un incident dure
CHAIN_GAP_ALERT_MIN = 10     # trou dans la chaîne au-delà duquel on alerte

FR_MAX_AGE_DAYS = 7          # prix de station plus vieux → ignoré (station fermée, abandonnée)
FR_MIN_STATIONS = 500        # en dessous, la médiane n'est pas fiable → échec de la source
FR_MIN_STATIONS_AUTOROUTE = 100  # beaucoup moins de stations d'autoroute que de stations sur route
FR_FAIL_THRESHOLD = 3        # la France n'est interrogée qu'à chaque notif → 3 échecs d'affilée

FUELS = ("diesel", "sp95", "sp98", "mazout")
CARBURANTS = ("diesel", "sp95", "sp98")
FUEL_LABELS = {"diesel": "Diesel", "sp95": "SP95", "sp98": "SP98", "mazout": "Mazout",
               "diesel_a": "Diesel autoroute", "sp95_a": "SP95 autoroute", "sp98_a": "SP98 autoroute"}

SOURCES = {
    "carbu_lu": {
        "label": "carbu.com (Luxembourg)",
        "url": "https://carbu.com/luxembourg/index.php/prixmaximum",
        "country": "lu",
    },
    "petrol_lu": {
        "label": "petrol.lu (mazout Luxembourg)",
        "url": "https://www.petrol.lu/prix-officiels/",
        "country": "lu",
    },
    "spf_be": {
        "label": "SPF Économie (Belgique)",
        "url": "https://petrolprices.economie.fgov.be/petrolprices/?locale=fr",
        "country": "be",
    },
    "fr_stations": {
        "label": "prix-carburants.gouv.fr (France)",
        "url": "https://data.economie.gouv.fr/api/explore/v2.1/catalog/datasets/"
               "prix-des-carburants-en-france-flux-instantane-v2/exports/json",
        "params": {"select": "pop,gazole_prix,gazole_maj,gazole_rupture_type,"
                             "e10_prix,e10_maj,e10_rupture_type,"
                             "sp98_prix,sp98_maj,sp98_rupture_type"},
        "timeout": (10, 90),
        "country": "fr",
        "display_only": True,
    },
    "fr_fioul": {
        "label": "Bulletin pétrolier UE (mazout France)",
        "url": "https://energy.ec.europa.eu/document/download/"
               "264c2d0f-f161-4ea3-a777-78faae59bea0_en?filename="
               "Weekly%20Oil%20Bulletin%20Weekly%20prices%20with%20Taxes%20-%202024-02-19.xlsx",
        "country": "fr",
        "display_only": True,
    },
}

MAIN_SOURCES = ("carbu_lu", "petrol_lu", "spf_be")   # déclenchent les notifs
FR_SOURCES = ("fr_stations", "fr_fioul")             # affichage seulement


class SourceError(Exception):
    """Échec d'une source. immediate=True → alerte sans attendre 1 h."""

    def __init__(self, message, immediate=False):
        super().__init__(message)
        self.immediate = immediate


# ─── Utilitaires ────────────────────────────────────────────────────
def now_local():
    return datetime.now(TZ)


def log(msg):
    print(f"[{now_local():%H:%M:%S}] {msg}", flush=True)


PRICE_RE = re.compile(r"(\d+)[.,](\d+)")
DATE_RE = re.compile(r"(\d{1,2})[./](\d{1,2})[./](\d{4})")


def parse_price(text):
    m = PRICE_RE.search(text or "")
    return Decimal(f"{m.group(1)}.{m.group(2)}") if m else None


def parse_date(text):
    m = DATE_RE.search(text or "")
    if not m:
        return None
    try:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def check_range(source, fuel, price):
    if not (PRICE_MIN <= price <= PRICE_MAX):
        raise SourceError(
            f"Prix {FUEL_LABELS[fuel]} hors fourchette sur {SOURCES[source]['label']} : "
            f"{price} € (attendu entre {PRICE_MIN} € et {PRICE_MAX} €)",
            immediate=True,
        )
    return price


def trunc3(value):
    """Tronque à 3 décimales, comme l'affichage en station (jamais au-dessus du maximum)."""
    return value.quantize(Decimal("0.001"), rounding=ROUND_DOWN)


def fmt_price(value):
    return f"{trunc3(value):.3f}€"


def fmt_date(d):
    return d.strftime("%d/%m") if d else "?"


def fmt_dt(value):
    if isinstance(value, str):
        value = datetime.fromisoformat(value)
    return value.astimezone(TZ).strftime("%d/%m/%Y à %H:%M")


def minutes_between(start, end):
    return int((end - start).total_seconds() // 60)


def cell_text(cell):
    return cell.get_text(" ", strip=True).replace("\xa0", " ")


def own_rows(table):
    """Lignes du tableau lui-même, sans celles d'éventuels tableaux imbriqués."""
    return [r for r in table.find_all("tr") if r.find_parent("table") is table]


# ─── Récupération ───────────────────────────────────────────────────
def fetch(url, timeout=REQUEST_TIMEOUT, params=None):
    last_error = None
    for attempt in range(1, REQUEST_ATTEMPTS + 1):
        try:
            r = requests.get(url, headers=HEADERS, timeout=timeout, params=params)
            if r.status_code == 400 and params:
                # Paramètre refusé par l'API : on retente sans filtre plutôt que d'échouer.
                log(f"  ⚠️ Paramètres refusés par {url} — nouvelle tentative sans filtre.")
                params = None
                r = requests.get(url, headers=HEADERS, timeout=timeout)
            r.raise_for_status()
            return r.content
        except requests.RequestException as e:
            last_error = e
            log(f"  ⏳ Tentative {attempt}/{REQUEST_ATTEMPTS} échouée sur {url} : {e}")
            if attempt < REQUEST_ATTEMPTS:
                time.sleep(RETRY_DELAY)
    raise SourceError(f"Site injoignable : {last_error}")


def parse_carbu_lu(html):
    """Tableau des prix maximum : une colonne par date (aujourd'hui, et demain si annoncé).
    On prend pour chaque carburant le prix le plus à droite, avec la date de sa colonne."""
    soup = BeautifulSoup(html, "html.parser")
    wanted = {"diesel": "diesel", "super 95": "sp95", "super 98": "sp98"}

    for table in soup.find_all("table"):
        rows = own_rows(table)
        if len(rows) < 2:
            continue
        header = rows[0].find_all(["th", "td"])
        date_cols = [d for d in (parse_date(cell_text(c)) for c in header[1:]) if d]
        if not date_cols:
            continue  # pas le tableau daté (ex. tableau « Évolution »)

        found = {}
        for row in rows[1:]:
            cells = row.find_all(["td", "th"])
            if not cells:
                continue
            label = cell_text(cells[0]).lower()
            if label not in wanted:
                continue
            # On ne garde que les cellules de prix (ou « n.d. ») : les cellules vides
            # ou décoratives ne décalent plus la lecture.
            values = [cell_text(c) for c in cells[1:]]
            values = [v for v in values if parse_price(v) is not None or "n.d" in v.lower()]
            # Alignement depuis la droite : dernière valeur ↔ dernière date.
            for k in range(1, min(len(values), len(date_cols)) + 1):
                price = parse_price(values[-k])
                if price is not None:
                    found[wanted[label]] = (price, date_cols[-k])
                    break

        if set(found) == set(wanted.values()):
            prices = {f: check_range("carbu_lu", f, p) for f, (p, _) in found.items()}
            dates = {f: d for f, (_, d) in found.items()}
            return prices, dates

    raise SourceError("Tableau des prix maximum introuvable ou incomplet sur carbu.com")


def parse_petrol_lu(html):
    """Tableau « Prix maxima » : on lit la ligne TVAC la plus récente,
    colonne « Gasoil chauff. BTS » repérée par son en-tête."""
    soup = BeautifulSoup(html, "html.parser")
    container = soup.find("div", class_="prices-table")
    table = container.find("table") if container else soup.find("table")
    if table is None:
        raise SourceError("Tableau des prix introuvable sur petrol.lu")

    rows = own_rows(table)
    col = None
    for row in rows:
        cells = row.find_all(["th", "td"])
        if cells and "date" in cell_text(cells[0]).lower():
            for i, c in enumerate(cells):
                if "bts" in cell_text(c).lower():
                    col = i
            break
    if col is None:
        raise SourceError("Colonne « Gasoil chauff. BTS » introuvable sur petrol.lu")

    best = None
    for row in rows:
        cells = row.find_all(["td", "th"])
        if len(cells) <= col or cell_text(cells[-1]).upper() != "TVAC":
            continue
        price = parse_price(cell_text(cells[col]))
        if price is None:
            continue
        d = parse_date(cell_text(cells[0]))
        if best is None or (d and (best[0] is None or d > best[0])):
            best = (d, price)

    if best is None:
        raise SourceError("Ligne TVAC du mazout introuvable sur petrol.lu")
    d, price = best
    return {"mazout": check_range("petrol_lu", "mazout", price)}, {"mazout": d}


def parse_spf_be(html):
    """Tarif officiel : « Produits pétroliers : tarif valable à partir du JJ.MM.AAAA »."""
    soup = BeautifulSoup(html, "html.parser")
    text = " ".join(soup.get_text(" ").split())
    m = re.search(r"valable\s+(?:à|a)\s+partir\s+du\s+(\d{1,2}[./]\d{1,2}[./]\d{4})",
                  text, re.IGNORECASE)
    valid_from = parse_date(m.group(1)) if m else None

    rows = soup.select("tbody#petrolTable_data tr") or soup.find_all("tr")
    prices = {}
    for row in rows:
        cells = row.find_all("td")
        if len(cells) < 2:
            continue
        label = cell_text(cells[0]).lower()
        price = parse_price(cell_text(cells[1]))
        if price is None:
            continue
        fuel = None
        if "95" in label and "e10" in label:
            fuel = "sp95"
        elif "98" in label and "e5" in label:
            fuel = "sp98"
        elif "diesel" in label and "b7" in label:
            fuel = "diesel"
        elif "chauffage" in label and "partir de 2000" in label:
            fuel = "mazout"
        if fuel:
            prices[fuel] = check_range("spf_be", fuel, price)

    missing = [FUEL_LABELS[f] for f in FUELS if f not in prices]
    if missing:
        raise SourceError(f"Prix manquants sur le site du SPF Économie : {', '.join(missing)}")
    return prices, {f: valid_from for f in FUELS}


def parse_fr_stations(raw):
    """Flux officiel des ~10 000 stations françaises : médiane des prix des stations
    hors autoroute (pop = "R") et, séparément, des stations d'autoroute (pop = "A").
    Même méthode pour les deux : prix de moins de FR_MAX_AGE_DAYS jours, ruptures et
    valeurs aberrantes exclues. Clés : diesel/sp95/sp98 (route), diesel_a/sp95_a/sp98_a."""
    try:
        data = json.loads(raw)
    except ValueError as e:
        raise SourceError(f"Réponse illisible de prix-carburants.gouv.fr : {e}")
    if isinstance(data, dict):
        data = data.get("results", [])
    if not isinstance(data, list):
        raise SourceError("Format inattendu des données de prix-carburants.gouv.fr")

    keys = {"diesel": "gazole", "sp95": "e10", "sp98": "sp98"}
    cutoff = datetime.now(timezone.utc) - timedelta(days=FR_MAX_AGE_DAYS)
    values = {fuel + suffix: [] for suffix in ("", "_a") for fuel in keys}
    for station in data:
        if not isinstance(station, dict) or station.get("pop") not in ("R", "A"):
            continue
        suffix = "" if station.get("pop") == "R" else "_a"
        for fuel, key in keys.items():
            price, updated = station.get(f"{key}_prix"), station.get(f"{key}_maj")
            if price is None or not updated:
                continue
            if station.get(f"{key}_rupture_type") == "temporaire":
                continue
            try:
                when = datetime.fromisoformat(str(updated))
                price = Decimal(str(price))
            except (ValueError, ArithmeticError):
                continue
            if when.tzinfo is None:
                when = when.replace(tzinfo=timezone.utc)
            if when < cutoff or not (PRICE_MIN <= price <= PRICE_MAX):
                continue  # prix périmé ou aberrant d'une station isolée : ignoré
            values[fuel + suffix].append(price)

    prices = {}
    for fuel, vals in values.items():
        minimum = FR_MIN_STATIONS_AUTOROUTE if fuel.endswith("_a") else FR_MIN_STATIONS
        if len(vals) < minimum:
            raise SourceError(f"Trop peu de stations à jour pour le {FUEL_LABELS[fuel]} "
                              f"sur prix-carburants.gouv.fr ({len(vals)}, minimum {minimum})")
        prices[fuel] = check_range("fr_stations", fuel, statistics.median(vals))
    log("  🇫🇷 Stations retenues : "
        + ", ".join(f"{FUEL_LABELS[f]} {len(v)}" for f, v in values.items()))
    today = now_local().date()
    return prices, {fuel: today for fuel in prices}


def parse_fr_fioul(raw):
    """Bulletin pétrolier hebdomadaire de la Commission européenne (prix avec taxes) :
    ligne « France », colonne « Gas oil de chauffage », date du relevé en tête de fichier."""
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise SourceError(f"Fichier du Bulletin pétrolier illisible : {e!r}")
    rows = list(workbook.worksheets[0].iter_rows(values_only=True))

    col = None
    for row in rows[:5]:
        for i, cell in enumerate(row):
            text = str(cell).lower() if cell is not None else ""
            if "chauffage" in text or "heating gas oil" in text:
                col = i
                break
        if col is not None:
            break
    if col is None:
        raise SourceError("Colonne « Gas oil de chauffage » introuvable dans le Bulletin pétrolier")

    bulletin_date, divisor = None, Decimal(1)
    for row in rows[:5]:
        for cell in row:
            if isinstance(cell, datetime) and bulletin_date is None:
                bulletin_date = cell.date()
        unit = row[col] if len(row) > col else None
        if isinstance(unit, str) and unit.strip().lower().replace(" ", "") == "1000l":
            divisor = Decimal(1000)

    for row in rows:
        if row and isinstance(row[0], str) and row[0].strip().lower() == "france":
            value = row[col] if len(row) > col else None
            if value is None:
                raise SourceError("Prix du fioul France absent du Bulletin pétrolier")
            price = Decimal(str(value)) / divisor
            return ({"mazout": check_range("fr_fioul", "mazout", price)},
                    {"mazout": bulletin_date})
    raise SourceError("Ligne « France » introuvable dans le Bulletin pétrolier")


PARSERS = {"carbu_lu": parse_carbu_lu, "petrol_lu": parse_petrol_lu, "spf_be": parse_spf_be,
           "fr_stations": parse_fr_stations, "fr_fioul": parse_fr_fioul}


def fetch_source(name):
    source = SOURCES[name]
    try:
        raw = fetch(source["url"], timeout=source.get("timeout", REQUEST_TIMEOUT),
                    params=source.get("params"))
        return PARSERS[name](raw)
    except SourceError as e:
        return e
    except Exception as e:  # bug de parsing inattendu : traité comme un échec de source
        return SourceError(f"Erreur inattendue : {e!r}")


def fetch_all(names):
    with ThreadPoolExecutor(max_workers=len(names)) as pool:
        futures = {name: pool.submit(fetch_source, name) for name in names}
        return {name: f.result() for name, f in futures.items()}


# ─── Cache ──────────────────────────────────────────────────────────
def empty_cache():
    return {"version": 2, "prices": {"lu": {}, "be": {}, "fr": {}}, "health": {}, "incidents": {},
            "fioul_watch": {}}


def load_cache():
    if not CACHE_FILE.exists():
        return empty_cache()
    try:
        data = json.loads(CACHE_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        log(f"⚠️ Cache illisible ({e}) — on repart de zéro.")
        return empty_cache()

    if data.get("version") == 2:
        for key, value in empty_cache().items():
            data.setdefault(key, value)
        for country in ("lu", "be", "fr"):
            data["prices"].setdefault(country, {})
        return data

    # Ancien format (version actuellement en prod) : {"lu": {...}, "be": {...}}
    cache = empty_cache()
    for country in ("lu", "be"):
        for fuel, value in (data.get(country) or {}).items():
            if fuel in FUELS and isinstance(value, (int, float)):
                cache["prices"][country][fuel] = {"value": str(Decimal(str(value))), "date": None}
    log("Ancien cache converti au nouveau format.")
    return cache


def save_cache(cache):
    tmp = CACHE_FILE.with_suffix(".tmp")
    tmp.write_text(json.dumps(cache, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
                   encoding="utf-8")
    os.replace(tmp, CACHE_FILE)


def append_history(rows):
    """Journal des changements détectés : sert à mesurer les heures de publication."""
    if not rows:
        return
    new_file = not HISTORY_FILE.exists()
    with HISTORY_FILE.open("a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if new_file:
            writer.writerow(["detecte_le", "source", "carburant",
                             "ancien_prix", "nouveau_prix", "valable_des"])
        writer.writerows(rows)


# ─── Envois ─────────────────────────────────────────────────────────
def send_ntfy(url, title, message, priority="high", tags=""):
    """Retourne (succès, erreur)."""
    if not url:
        log(f"⚠️ Canal ntfy non configuré — message non envoyé :\n{title}\n{message}")
        return True, None
    try:
        r = requests.post(
            url,
            data=message.encode("utf-8"),
            headers={"Title": title.encode("utf-8"), "Priority": priority, "Tags": tags},
            timeout=15,
        )
        r.raise_for_status()
        return True, None
    except requests.RequestException as e:
        log(f"❌ Envoi ntfy échoué : {e}")
        return False, str(e)


def send_mail(subject, body):
    if not (MAIL_ADDRESS and MAIL_APP_PASSWORD):
        log("⚠️ MAIL_ADDRESS ou MAIL_APP_PASSWORD non configuré — mail non envoyé.")
        return True
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = MAIL_ADDRESS
    msg["To"] = MAIL_ADDRESS
    msg.set_content(body)
    try:
        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=30,
                              context=ssl.create_default_context()) as server:
            server.login(MAIL_ADDRESS, MAIL_APP_PASSWORD)
            server.send_message(msg)
        log(f"📧 Mail envoyé : {subject}")
        return True
    except Exception as e:
        log(f"❌ Envoi du mail échoué : {e!r}")
        return False


def send_alert_once(title, body):
    """Alerte ponctuelle (pas d'incident suivi) : mail + canal technique."""
    send_mail(f"URGENT !!!! Fuel Alert Bot — {title}", body)
    send_ntfy(NTFY_ALERT_URL, f"🔧 {title}", body, priority="urgent", tags="warning")


class Alerts:
    """Incidents suivis dans le cache : une alerte à l'ouverture, un rappel toutes
    les 24 h, un message « Résolu » à la fermeture."""

    def __init__(self, cache, now):
        self.incidents = cache["incidents"]
        self.now = now
        self.resolved = []

    def open(self, key, title, body):
        if key not in self.incidents:
            self.incidents[key] = {
                "title": title, "body": body,
                "opened": self.now.isoformat(timespec="seconds"),
                "mail_sent": False, "ntfy_sent": False,
                "last_sent": None, "reminder": False,
            }
            log(f"🚨 Incident ouvert : {title}")

    def resolve(self, key):
        incident = self.incidents.pop(key, None)
        if incident:
            self.resolved.append(incident)
            log(f"✅ Incident résolu : {incident['title']}")

    def any_open(self):
        return bool(self.incidents)

    def flush(self):
        for inc in self.incidents.values():
            if inc["mail_sent"] and inc["ntfy_sent"] and inc["last_sent"]:
                if self.now - datetime.fromisoformat(inc["last_sent"]) >= timedelta(hours=REMINDER_HOURS):
                    inc.update(mail_sent=False, ntfy_sent=False, reminder=True)

            prefix = "(Rappel) " if inc["reminder"] else ""
            body = (f"{inc['body']}\n\nIncident ouvert le {fmt_dt(inc['opened'])}.\n"
                    "Le bot continue de tourner et t'enverra un message « Résolu » "
                    "dès que le problème disparaît.")
            sent_now = False
            if not inc["mail_sent"]:
                inc["mail_sent"] = send_mail(f"URGENT !!!! Fuel Alert Bot — {prefix}{inc['title']}", body)
                sent_now = sent_now or inc["mail_sent"]
            if not inc["ntfy_sent"]:
                inc["ntfy_sent"], _ = send_ntfy(NTFY_ALERT_URL, f"🔧 {prefix}{inc['title']}", body,
                                                priority="urgent", tags="warning")
                sent_now = sent_now or inc["ntfy_sent"]
            if sent_now and inc["mail_sent"] and inc["ntfy_sent"]:
                inc["last_sent"] = self.now.isoformat(timespec="seconds")

        for inc in self.resolved:
            body = (f"Le problème suivant est résolu :\n\n{inc['title']}\n\n"
                    f"Ouvert le {fmt_dt(inc['opened'])}, résolu le {fmt_dt(self.now)}.")
            send_mail(f"✅ Résolu — Fuel Alert Bot — {inc['title']}", body)
            send_ntfy(NTFY_ALERT_URL, f"✅ Résolu : {inc['title']}", body,
                      priority="default", tags="white_check_mark")


# ─── Santé des sources ──────────────────────────────────────────────
def new_health():
    return {"fails": 0, "since": None, "last_error": None}


def update_source_health(cache, alerts, name, result, now):
    health = cache["health"].setdefault(name, new_health())
    source = SOURCES[name]

    if isinstance(result, SourceError):
        health["fails"] += 1
        health["since"] = health["since"] or now.isoformat(timespec="seconds")
        health["last_error"] = str(result)
        log(f"  ❌ {source['label']} : {result} (échec n°{health['fails']})")
        threshold = FR_FAIL_THRESHOLD if source.get("display_only") else FAIL_THRESHOLD
        if result.immediate or health["fails"] >= threshold:
            since = datetime.fromisoformat(health["since"])
            if result.immediate:
                title = f"Prix aberrant sur {source['label']}"
            elif source.get("display_only"):
                title = f"{source['label']} en échec ({health['fails']} tentatives d'affilée)"
            else:
                title = f"{source['label']} en échec depuis {minutes_between(since, now)} min"
            body = (f"Source : {source['label']}\n"
                    f"Adresse : {source['url']}\n"
                    f"Premier échec : {fmt_dt(since)}\n"
                    f"Échecs consécutifs : {health['fails']}\n"
                    f"Dernière erreur : {result}\n\n"
                    "Tant que ce problème dure, les prix de cette source ne sont plus mis à jour. "
                    "Les autres sources continuent de fonctionner normalement.")
            alerts.open(f"source:{name}", title, body)
        return

    if health["fails"]:
        log(f"  ✅ {source['label']} répond de nouveau après {health['fails']} échec(s).")
    cache["health"][name] = new_health()
    alerts.resolve(f"source:{name}")

    _, dates = result
    missing = [FUEL_LABELS[f] for f, d in dates.items() if d is None]
    if missing:
        alerts.open(
            f"date:{name}",
            f"Date de validité illisible sur {source['label']}",
            f"Les prix de {source['label']} sont lus correctement, mais pas leur date de "
            f"validité ({', '.join(missing)}). La notification affichera « ? » à la place "
            f"de la date.\nAdresse : {source['url']}",
        )
    else:
        alerts.resolve(f"date:{name}")


def update_ntfy_health(cache, alerts, ok, error, now):
    health = cache["health"].setdefault("ntfy", new_health())
    if ok:
        cache["health"]["ntfy"] = new_health()
        alerts.resolve("ntfy")
        return
    health["fails"] += 1
    health["since"] = health["since"] or now.isoformat(timespec="seconds")
    health["last_error"] = error
    if health["fails"] >= NTFY_FAIL_THRESHOLD:
        since = datetime.fromisoformat(health["since"])
        alerts.open(
            "ntfy",
            "Les notifications de prix ne partent plus",
            f"ntfy refuse les notifications de prix depuis {fmt_dt(since)} "
            f"({health['fails']} tentatives).\nDernière erreur : {error}\n\n"
            "Les changements de prix détectés seront renvoyés automatiquement "
            "dès que ntfy les accepte de nouveau.",
        )


# ─── Notification de prix ───────────────────────────────────────────
def date_line(flag, entries):
    def d(fuel):
        entry = entries.get(fuel)
        return entry["date"] if entry else None

    carb_dates = [d(f) for f in CARBURANTS]
    if all(x == carb_dates[0] for x in carb_dates):
        carb = f"Carburants dès le {fmt_date(carb_dates[0])}"
    else:
        carb = ", ".join(f"{FUEL_LABELS[f]} dès le {fmt_date(d(f))}" for f in CARBURANTS)
    return f"📅 {flag} {carb} · Mazout dès le {fmt_date(d('mazout'))}"


def fr_date_line(entries, fioul_published=None):
    def d(fuel):
        entry = entries.get(fuel)
        return entry["date"] if entry else None
    mazout = f"Mazout : prix du {fmt_date(d('mazout'))}"
    if fioul_published:
        mazout += f", publié le {fmt_date(fioul_published)}"
    return f"📅 🇫🇷 Carburants relevés le {fmt_date(d('diesel'))} · {mazout}"


def fr_lines(fuel, current_fr, old_fr):
    """Lignes France : hausse/baisse calculée depuis la notif précédente, comme LU/BE."""
    def line(key, suffix):
        entry, old = current_fr.get(key), old_fr.get(key)
        changed = entry is not None and old is not None and trunc3(entry["value"]) != trunc3(old)
        return f"{price_line('🇫🇷', entry, old, changed)} {suffix}"
    if fuel == "mazout":
        return [line("mazout", "moyenne nationale")]
    return [line(fuel, "hors autoroute"), line(fuel + "_a", "sur autoroute 🛣️")]


def trophy(entries):
    """Pays le moins cher parmi ceux qui ont un prix (comparaison à 3 décimales)."""
    names = {"lu": "🏆 Go Lux. !", "be": "🏆 Go Belgique !", "fr": "🏆 Go France !"}
    prices = {c: trunc3(e["value"]) for c, e in entries.items() if e}
    if len(prices) < 2:
        return None
    best = min(prices.values())
    winners = [c for c, v in prices.items() if v == best]
    return names[winners[0]] if len(winners) == 1 else "🤝 Égalité !"


def price_line(flag, entry, old, changed):
    if entry is None:
        return f"{flag} n/d"
    value = entry["value"]
    if old is None:
        status = "🆕 Nouveau"
    elif not changed:
        status = "⚖️ Stable"
    else:
        delta = trunc3(value) - trunc3(old)
        status = f"📈 Hausse de +{delta:.3f}€" if delta > 0 else f"📉 Baisse de {abs(delta):.3f}€"
    return f"{flag} {fmt_price(value)} ({status})"


def build_notification(current, old_values, changed, fioul_published=None):
    lu_changed = any(country == "lu" for country, _ in changed)
    be_changed = any(country == "be" for country, _ in changed)
    if lu_changed and be_changed:
        subtitle = "🇱🇺🇧🇪 Les deux pays changent leurs prix !"
    elif lu_changed:
        subtitle = "🇱🇺 Le Luxembourg change ses prix !"
    else:
        subtitle = "🇧🇪 La Belgique change ses prix !"

    lines = [subtitle, "",
             date_line("🇱🇺", current["lu"]),
             date_line("🇧🇪", current["be"]),
             fr_date_line(current["fr"], fioul_published),
             ""]

    for fuel, emoji, label in (("diesel", "🟡", "DIESEL"), ("sp95", "🟢", "SP95"),
                               ("sp98", "🔵", "SP98"), ("mazout", "🔥", "MAZOUT CHAUFFAGE")):
        entries = {c: current[c].get(fuel) for c in ("lu", "be", "fr")}
        winner = trophy(entries)
        lines.append(f"{emoji} {label} | {winner}" if winner else f"{emoji} {label}")
        lines.append(price_line("🇱🇺", entries["lu"], old_values["lu"].get(fuel), ("lu", fuel) in changed))
        lines.append(price_line("🇧🇪", entries["be"], old_values["be"].get(fuel), ("be", fuel) in changed))
        lines.extend(fr_lines(fuel, current["fr"], old_values["fr"]))
        lines.append("")

    return "⛽️ Alerte Carburants", "\n".join(lines).strip()


# ─── Programme principal ────────────────────────────────────────────
def check_revived_chain(now):
    """Si le watchdog a dû relancer la chaîne, on prévient (sauf premier démarrage)."""
    if os.environ.get("REVIVED") != "true":
        return
    last_seen = os.environ.get("LAST_SEEN", "").strip()
    if not last_seen:
        return
    seen = datetime.fromisoformat(last_seen.replace("Z", "+00:00")).astimezone(TZ)
    gap = minutes_between(seen, now)
    if gap >= CHAIN_GAP_ALERT_MIN:
        send_alert_once(
            f"Chaîne des 5 min interrompue pendant {gap} min",
            f"Dernier run avant l'arrêt : {fmt_dt(seen)}.\n"
            f"Chaîne relancée automatiquement par le watchdog le {fmt_dt(now)}.\n"
            "Cause la plus probable : une panne GitHub Actions (githubstatus.com).\n"
            "Aucune action n'est nécessaire de ta part.",
        )


def main():
    now = now_local()
    log("Vérification des prix…")
    cache = load_cache()
    alerts = Alerts(cache, now)
    results = fetch_all(MAIN_SOURCES + ("fr_fioul",))
    fioul_result = results.pop("fr_fioul")

    for name, result in results.items():
        update_source_health(cache, alerts, name, result, now)

    # Surveillance du Bulletin pétrolier à chaque run (14 Ko), uniquement pour dater sa
    # publication. Un échec ici ne compte pas : seuls les échecs au moment des notifs comptent.
    watch = cache.setdefault("fioul_watch", {})
    if isinstance(fioul_result, SourceError):
        log(f"  ℹ️ Bulletin pétrolier non consulté ce run : {fioul_result}")
    else:
        bulletin_date = fioul_result[1].get("mazout")
        bulletin_iso = bulletin_date.isoformat() if bulletin_date else None
        if bulletin_iso and watch.get("bulletin_date") != bulletin_iso:
            first_seen = bool(watch)  # au tout premier passage, la date de publication est inconnue
            watch.clear()
            watch["bulletin_date"] = bulletin_iso
            watch["published"] = now.date().isoformat() if first_seen else None
            log(f"  🇫🇷 Nouveau Bulletin pétrolier (prix du {fmt_date(bulletin_date)})"
                + (f", publié le {fmt_date(now.date())}." if first_seen else " — date de publication inconnue."))

    # Valeurs connues (cache) puis remplacées par les valeurs fraîches des sources OK.
    current = {"lu": {}, "be": {}, "fr": {}}
    for country in ("lu", "be", "fr"):
        for fuel, entry in cache["prices"][country].items():
            current[country][fuel] = {
                "value": Decimal(entry["value"]),
                "date": date.fromisoformat(entry["date"]) if entry.get("date") else None,
            }
    old_values = {c: {f: e["value"] for f, e in current[c].items()} for c in current}

    changed, history_rows = [], []
    for name, result in results.items():
        if isinstance(result, SourceError):
            continue
        prices, dates = result
        country = SOURCES[name]["country"]
        log(f"  ✅ {SOURCES[name]['label']} : "
            + ", ".join(f"{FUEL_LABELS[f]} {fmt_price(v)} (dès le {fmt_date(dates.get(f))})"
                        for f, v in prices.items()))
        for fuel, value in prices.items():
            previous = old_values[country].get(fuel)
            if previous is None or trunc3(previous) != trunc3(value):
                changed.append((country, fuel))
                history_rows.append([
                    now.isoformat(timespec="seconds"), name, fuel,
                    "" if previous is None else str(previous), str(value),
                    dates[fuel].isoformat() if dates.get(fuel) else "",
                ])
            current[country][fuel] = {"value": value, "date": dates.get(fuel)}

    commit_prices = True
    if changed:
        log(f"  🚨 Changement détecté : "
            + ", ".join(f"{c.upper()} {FUEL_LABELS[f]}" for c, f in changed))
        # La France n'est interrogée qu'au moment d'envoyer une notif (affichage seulement).
        fr_results = fetch_all(("fr_stations",))
        fr_results["fr_fioul"] = fioul_result
        for name, result in fr_results.items():
            update_source_health(cache, alerts, name, result, now)
            if isinstance(result, SourceError):
                log("  🇫🇷 Valeurs précédentes conservées pour cette source.")
                continue
            prices, dates = result
            log(f"  ✅ {SOURCES[name]['label']} : "
                + ", ".join(f"{FUEL_LABELS[f]} {fmt_price(v)}" for f, v in prices.items()))
            for fuel, value in prices.items():
                current["fr"][fuel] = {"value": value, "date": dates.get(fuel)}
        mazout_fr = current["fr"].get("mazout")
        fioul_published = None
        if (mazout_fr and mazout_fr["date"] and watch.get("published")
                and watch.get("bulletin_date") == mazout_fr["date"].isoformat()):
            fioul_published = date.fromisoformat(watch["published"])
        title, message = build_notification(current, old_values, changed, fioul_published)
        ok, error = send_ntfy(NTFY_URL, title, message, priority="high",
                              tags="fuelpump,rotating_light")
        update_ntfy_health(cache, alerts, ok, error, now)
        if ok:
            log(f"  ✅ Notification envoyée : {title}")
        else:
            commit_prices = False  # on retentera au prochain run
            log("  Prix NON enregistrés — la notification sera renvoyée au prochain run.")
    else:
        log("  Aucun changement de prix.")

    if commit_prices:
        for country in ("lu", "be", "fr"):
            cache["prices"][country] = {
                fuel: {"value": str(e["value"]), "date": e["date"].isoformat() if e["date"] else None}
                for fuel, e in current[country].items()
            }
        append_history(history_rows)

    check_revived_chain(now)
    alerts.flush()
    save_cache(cache)

    if alerts.any_open():
        log(f"⚠️ Incident(s) en cours : {', '.join(i['title'] for i in alerts.incidents.values())}")
        return 1
    return 0


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--chain-error":
        send_alert_once(
            "La chaîne des 5 min n'a pas pu lancer le run suivant",
            "GitHub a refusé trois fois de démarrer le run suivant.\n"
            "Le watchdog relancera la chaîne automatiquement (délai : environ 10 min, "
            "selon GitHub). Si cette alerte revient souvent, vérifie githubstatus.com.",
        )
        sys.exit(0)
    sys.exit(main())
